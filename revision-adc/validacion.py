#!/usr/bin/env python3
"""Analiza la codificación a texto completo de la muestra de validación.

Dos pasos:

  acuerdo   Compara las planillas de los dos codificadores: porcentaje de
            acuerdo y kappa de Cohen por variable. Escribe una planilla de
            consenso con los acuerdos ya puestos y los desacuerdos en blanco,
            en amarillo, con lo que puso cada uno en un comentario.
  contraste Compara la planilla de consenso, ya resuelta, con los códigos de
            los diccionarios (03-codificacion.csv): para cada categoría,
            precisión y sensibilidad de los diccionarios y kappa. Resume además
            la elegibilidad, la evaluación de resultados y las tres dimensiones.

Uso: validacion.py acuerdo A.xlsx B.xlsx consenso.xlsx 12-acuerdo.json
     validacion.py contraste consenso.xlsx 03-codificacion.csv 13-validacion.json
"""
import csv
import json
import sys
from collections import Counter

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from planillas_validacion import CATEGORIAS, DIMENSIONES

AMARILLO = PatternFill("solid", fgColor="FFFF00")
CODIGOS = ["Texto completo leído", "Elegible", "Motivo de exclusión"] + CATEGORIAS + ["Evaluación de resultados"] + DIMENSIONES
PERIODOS = [("2016–2019", 2016, 2019), ("2020–2023", 2020, 2023), ("2024–2026", 2024, 2026)]


def leer(ruta):
    ws = load_workbook(ruta)["Codificación"]
    cab = [c.value for c in ws[1]]
    filas = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        x = dict(zip(cab, fila))
        if x.get("EID"):
            filas[x["EID"]] = {k: (str(v).strip() if v is not None else "") for k, v in x.items()}
    return filas


def kappa(pares):
    """Kappa de Cohen para una lista de pares de etiquetas."""
    n = len(pares)
    if not n:
        return None, None
    po = sum(a == b for a, b in pares) / n
    ca, cb = Counter(a for a, _ in pares), Counter(b for _, b in pares)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return round(po, 3), (round((po - pe) / (1 - pe), 3) if pe < 1 else None)


def acuerdo(ruta_a, ruta_b, ruta_consenso, ruta_json):
    a, b = leer(ruta_a), leer(ruta_b)
    comunes = [e for e in a if e in b and a[e]["Elegible"] and b[e]["Elegible"]]
    resultado = {"estudios_codificados_por_ambos": len(comunes), "variables": {}}
    for v in CODIGOS:
        # El motivo solo se compara cuando los dos excluyen; el contenido, cuando los dos incluyen
        if v in ("Texto completo leído", "Elegible"):
            filas = comunes
        elif v == "Motivo de exclusión":
            filas = [e for e in comunes if a[e]["Elegible"] == b[e]["Elegible"] == "No"]
        else:
            filas = [e for e in comunes if a[e]["Elegible"] == b[e]["Elegible"] == "Sí"]
        pares = [(a[e][v] or "No", b[e][v] or "No") for e in filas if v in a[e]]
        po, k = kappa(pares)
        resultado["variables"][v] = {"n": len(pares), "acuerdo": po, "kappa": k}
    wb = load_workbook(ruta_a)
    ws = wb["Codificación"]
    cab = [c.value for c in ws[1]]
    for fila in ws.iter_rows(min_row=2):
        eid = fila[cab.index("EID")].value
        if eid not in b:
            continue
        for j, v in enumerate(cab):
            if v in CODIGOS and (a[eid][v] or "") != (b[eid][v] or ""):
                fila[j].value, fila[j].fill = None, AMARILLO
                fila[j].comment = Comment(f"A: {a[eid][v] or '(vacío)'} · B: {b[eid][v] or '(vacío)'}", "validacion.py")
            elif v and (v.startswith("Cita") or v in ("Notas", "Otra tradición (cuál)")) and b[eid].get(v):
                fila[j].value = f"A: {a[eid][v]} | B: {b[eid][v]}"
    wb["Instrucciones"]["A1"] = "Planilla de consenso: resolved juntos las celdas en amarillo (el comentario dice qué puso cada uno)."
    wb.save(ruta_consenso)
    json.dump(resultado, open(ruta_json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    for v, r in resultado["variables"].items():
        print(f"{v:62} n={r['n']:3} acuerdo={r['acuerdo']} kappa={r['kappa']}")


def contraste(ruta_consenso, ruta_cod, ruta_json):
    c = leer(ruta_consenso)
    cod = {x["eid"]: x for x in csv.DictReader(open(ruta_cod, encoding="utf-8-sig"))}
    leidos = [e for e in c if c[e]["Texto completo leído"] == "Sí"]
    pendientes = [e for e in leidos if not c[e]["Elegible"]]
    if pendientes:
        sys.exit(f"Faltan decisiones de elegibilidad en {len(pendientes)} estudios: resuelve el consenso antes.")
    elegibles = [e for e in leidos if c[e]["Elegible"] == "Sí"]
    resultado = {
        "leidos": len(leidos), "elegibles": len(elegibles),
        "no_elegibles": dict(Counter(c[e]["Motivo de exclusión"] for e in leidos if c[e]["Elegible"] == "No")),
        "categorias": {}, "evaluacion": dict(Counter(c[e]["Evaluación de resultados"] or "(vacío)" for e in elegibles)),
    }
    for v in CATEGORIAS:
        vp = sum(c[e][v] == "Sí" and cod[e][v] == "1" for e in elegibles)
        fp = sum(c[e][v] != "Sí" and cod[e][v] == "1" for e in elegibles)
        fn = sum(c[e][v] == "Sí" and cod[e][v] != "1" for e in elegibles)
        po, k = kappa([(c[e][v] == "Sí", cod[e][v] == "1") for e in elegibles])
        resultado["categorias"][v] = {
            "texto_completo": vp + fn, "diccionario": vp + fp, "vp": vp, "fp": fp, "fn": fn,
            "precision": round(vp / (vp + fp), 3) if vp + fp else None,
            "sensibilidad": round(vp / (vp + fn), 3) if vp + fn else None, "acuerdo": po, "kappa": k}
    # Tres dimensiones: frecuencia, combinaciones y reparto por periodo
    def periodo(e):
        return next(p for p, i, f in PERIODOS if i <= int(cod[e]["anio"]) <= f)
    sigla = {"dimensión: Textual": "Txt", "dimensión: Infraestructural": "Inf", "dimensión: Transformadora": "Tra"}
    marca = {e: "+".join(sigla[d] for d in DIMENSIONES if c[e][d] == "Sí") or "ninguna" for e in elegibles}
    resultado["dimensiones"] = {
        "frecuencia": {d: sum(c[e][d] == "Sí" for e in elegibles) for d in DIMENSIONES},
        "combinaciones": dict(Counter(marca.values()).most_common()),
        "por_periodo": {p: {"n": sum(periodo(e) == p for e in elegibles),
                            **{d: sum(c[e][d] == "Sí" and periodo(e) == p for e in elegibles) for d in DIMENSIONES}}
                        for p, _, _ in PERIODOS},
    }
    json.dump(resultado, open(ruta_json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"Leídos {len(leidos)}, elegibles {len(elegibles)}")
    for v, r in resultado["categorias"].items():
        print(f"{v:62} precisión={r['precision']} sensibilidad={r['sensibilidad']} kappa={r['kappa']}")


if __name__ == "__main__":
    {"acuerdo": acuerdo, "contraste": contraste}[sys.argv[1]](*sys.argv[2:])
