#!/usr/bin/env python3
"""Prepara las planillas de codificación a texto completo, una por codificador.

Cada planilla lista los estudios de la muestra (10-muestra-validacion.csv) con
el archivo o el enlace del texto (11-textos-abiertos.csv) y deja en blanco las
variables que hay que codificar, con listas desplegables. No incluye los
códigos de los diccionarios, para que la lectura sea independiente de ellos.
Las definiciones están en manual-de-codificacion.md.

Los PDF descargados se copian a DIR_SALIDA/textos con un nombre legible
(número de fila, apellido y año), que es el que figura en la planilla. Los
textos que no se pudieron descargar llevan su enlace.

Uso: planillas_validacion.py 10-muestra-validacion.csv 11-textos-abiertos.csv DIR_PDF DIR_SALIDA
"""
import csv
import re
import shutil
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Mismas categorías y en el mismo orden que 03-codificacion.csv
CATEGORIAS = [
    "nivel: Infantil y primaria", "nivel: Secundaria", "nivel: Universidad", "nivel: Formación docente",
    "diseño: Revisión de literatura", "diseño: Teórico o conceptual", "diseño: Cualitativo",
    "diseño: Cuantitativo", "diseño: Diseño o intervención",
    "marco: Pedagogía crítica (Freire)", "marco: Multiliteracidades y nuevos estudios de literacidad",
    "marco: Alfabetización mediática", "marco: Ciudadanía digital",
    "marco: Perspectivas críticas de raza, género y decolonialidad",
    "tema: Inteligencia artificial", "tema: Desinformación y posverdad", "tema: Redes sociales",
    "tema: Algoritmos, datos y plataformas", "tema: Privacidad y vigilancia",
    "tema: Discurso de odio y radicalización", "tema: Lenguas extranjeras", "tema: Producción digital del alumnado",
]
DIMENSIONES = ["dimensión: Textual", "dimensión: Infraestructural", "dimensión: Transformadora"]
EVALUACION = ["0 Ninguna", "1 Evidencia cualitativa", "2 Medida cuantitativa sin comparación", "3 Diseño comparativo"]
MOTIVOS = ["La ADC no es el foco", "Contexto no formal", "Tipo de documento", "Otro (ver notas)"]
DATOS = ["N", "EID", "Primer autor", "Año", "Título", "Revista", "Texto (archivo o enlace)"]
NEGRITA = Font(name="Arial", bold=True)
NORMAL = Font(name="Arial")
CABECERA = PatternFill("solid", fgColor="E7E6E6")
CODIGO = PatternFill("solid", fgColor="FFF2CC")


def columnas():
    cols = DATOS + ["Texto completo leído", "Elegible", "Motivo de exclusión"] + CATEGORIAS + ["Evaluación de resultados"]
    for d in DIMENSIONES:
        cols += [d, d.replace("dimensión: ", "Cita y página: ")]
    return cols + ["Otra tradición (cuál)", "Notas"]


def nombre_legible(n, x):
    apellido = unicodedata.normalize("NFKD", x["primer_autor"].rsplit(" ", 1)[0]).encode("ascii", "ignore").decode()
    return f"{n:02d}-{re.sub(r'[^A-Za-z]+', '', apellido)}-{x['anio']}.pdf"


def hoja_codificacion(ws, muestra, textos):
    cols = columnas()
    for j, c in enumerate(cols, 1):
        celda = ws.cell(1, j, c)
        celda.font, celda.alignment = NEGRITA, Alignment(wrap_text=True, vertical="top")
        celda.fill = CABECERA if j <= len(DATOS) else CODIGO
    for i, x in enumerate(muestra, 2):
        t = textos.get(x["eid"], {})
        texto = (nombre_legible(i - 1, x) if t.get("archivo") else
                 t.get("pagina_url") or t.get("pdf_url") or f"https://doi.org/{x['doi']}")
        for j, v in enumerate([i - 1, x["eid"], x["primer_autor"], int(x["anio"]), x["titulo"], x["revista"], texto], 1):
            ws.cell(i, j, v).font = NORMAL
            ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
    ultima = len(muestra) + 1
    listas = {"Sí/No": '"Sí,No"', "motivo": '"' + ",".join(MOTIVOS) + '"',
              "evaluacion": '"' + ",".join(EVALUACION) + '"'}
    for nombre, formula in listas.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        for j, c in enumerate(cols, 1):
            tipo = ("motivo" if c == "Motivo de exclusión" else "evaluacion" if c == "Evaluación de resultados"
                    else "Sí/No" if c in ["Texto completo leído", "Elegible"] + CATEGORIAS + DIMENSIONES else None)
            if tipo == nombre:
                dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{ultima}")
    anchos = {"N": 5, "EID": 20, "Primer autor": 16, "Año": 6, "Título": 50, "Revista": 24,
              "Texto (archivo o enlace)": 28, "Evaluación de resultados": 24, "Notas": 40, "Otra tradición (cuál)": 20}
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = anchos.get(c, 34 if c.startswith("Cita") else 13)
    ws.row_dimensions[1].height = 75
    ws.freeze_panes = "H2"


def main():
    filas = list(csv.DictReader(open(sys.argv[1], encoding="utf-8-sig")))
    textos = {t["eid"]: t for t in csv.DictReader(open(sys.argv[2], encoding="utf-8-sig"))}
    pdf, salida = Path(sys.argv[3]), Path(sys.argv[4])
    muestra = [x for x in filas if x["papel"] == "muestra"]
    reservas = [x for x in filas if x["papel"] == "reserva"]
    (salida / "textos").mkdir(parents=True, exist_ok=True)
    for n, x in enumerate(muestra, 1):
        if textos.get(x["eid"], {}).get("archivo"):
            shutil.copy(pdf / textos[x["eid"]]["archivo"], salida / "textos" / nombre_legible(n, x))
    for codificador in ("A", "B"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Instrucciones"
        lineas = [
            f"Planilla del codificador {codificador}. Nombre: ",
            "Lee el manual de codificación antes de empezar y tenlo a mano.",
            "Codifica cada estudio leyendo el texto completo. No mires la planilla del otro codificador ni los códigos del artículo.",
            "Usa las listas desplegables. Si un estudio no es elegible, codifica solo «Texto completo leído», «Elegible» y «Motivo».",
            "Si no puedes leer un texto de la muestra, márcalo con «Texto completo leído: No» y codifica la primera reserva del mismo periodo en una fila nueva (hoja Reservas).",
            "Cuando termines, guarda el archivo con tu nombre y envíalo.",
        ]
        for i, l in enumerate(lineas, 1):
            ws.cell(i, 1, l).font = NEGRITA if i == 1 else NORMAL
        ws.column_dimensions["A"].width = 120
        hoja_codificacion(wb.create_sheet("Codificación"), muestra, textos)
        wr = wb.create_sheet("Reservas")
        for j, c in enumerate(["Periodo", "Orden", "EID", "Primer autor", "Año", "Título", "Texto (archivo o enlace)"], 1):
            wr.cell(1, j, c).font = NEGRITA
        for i, x in enumerate(reservas, 2):
            t = textos.get(x["eid"], {})
            texto = t.get("archivo") or t.get("pagina_url") or f"https://doi.org/{x['doi']}"
            for j, v in enumerate([x["periodo"], int(x["orden"]), x["eid"], x["primer_autor"], int(x["anio"]),
                                   x["titulo"], texto], 1):
                wr.cell(i, j, v).font = NORMAL
        for j, w in enumerate([11, 7, 20, 16, 6, 60, 30], 1):
            wr.column_dimensions[get_column_letter(j)].width = w
        wb.active = 1
        wb.save(salida / f"planilla-codificador-{codificador}.xlsx")


if __name__ == "__main__":
    main()
