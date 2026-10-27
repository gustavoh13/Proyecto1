#!/usr/bin/env python3
"""Genera el archivo suplementario S5 (codificación a texto completo, en inglés).

Una hoja con la codificación de cada estudio de la muestra y de las reservas
usadas (con las citas y páginas que sostienen cada dimensión), otra con la
precisión y la sensibilidad de los diccionarios y otra con los resúmenes de
13-validacion.json. Las notas en inglés están en notas-en.json.

Uso: suplementos.py 10-muestra-validacion.csv 11-textos-abiertos.csv 03-codificacion.csv \
         codificacion-texto-completo.jsonl notas-en.json 13-validacion.json S5.xlsx
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from planillas_validacion import CATEGORIAS, DIMENSIONES  # noqa: E402

ETIQUETAS = {
    "niv": {"IP": "Early childhood and primary", "SEC": "Secondary", "UNI": "Higher education", "FD": "Teacher education"},
    "dis": {"REV": "Literature review", "TEO": "Conceptual or theoretical", "CUAL": "Qualitative",
            "CUAN": "Quantitative", "INT": "Intervention or designed teaching"},
    "trad": {"FRE": "Critical pedagogy (Freirean)", "MUL": "Multiliteracies and New Literacy Studies",
             "MED": "Media literacy and educommunication", "CIU": "Digital citizenship",
             "RGD": "Race, gender and decolonial perspectives"},
    "obj": {"IA": "Artificial intelligence", "DES": "Disinformation and post-truth", "RED": "Social media",
            "ALG": "Algorithms, data and platforms", "PRI": "Privacy and surveillance",
            "ODI": "Hate speech and radicalisation", "LEN": "Second or foreign language education",
            "PRO": "Student media production"},
}
# Mismo orden que CATEGORIAS
EN_CATEGORIAS = [e for k in ("niv", "dis", "trad", "obj") for e in ETIQUETAS[k].values()]
VARIABLE = ["Level"] * 4 + ["Design"] * 5 + ["Tradition"] * 5 + ["Object of critique"] * 8
EVALUACION = ["0 None", "1 Qualitative evidence", "2 Quantitative measure without comparison", "3 Comparative design"]
MOTIVO = {"La ADC no es el foco": "CDL is not a central focus"}
PERIODO = {"2016–2019": "2016–2019", "2020–2023": "2020–2023", "2024–2026": "2024–2026"}
VIA = {"acuerdo entre fases": "Both passes agreed",
       "discrepancia: el título nombra la ADC": "Disagreement; title names CDL",
       "discrepancia: título digital y dos indicios de resumen o palabras clave":
           "Disagreement; title refers to a digital technology and two of K, A, P"}
DIM_EN = {"T": "Textual", "I": "Infrastructural", "R": "Transformative"}
NEGRITA, NORMAL = Font(name="Arial", bold=True), Font(name="Arial")
CABECERA = PatternFill("solid", fgColor="E7E6E6")


def escribir(ws, filas, anchos):
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila, 1):
            celda = ws.cell(i, j, v)
            celda.font = NEGRITA if i == 1 else NORMAL
            celda.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1:
                celda.fill = CABECERA
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"


def main():
    muestra, textos, cod, jsonl, notas, val, salida = sys.argv[1:8]
    filas = list(csv.DictReader(open(muestra, encoding="utf-8-sig")))
    textos = {t["eid"]: t for t in csv.DictReader(open(textos, encoding="utf-8-sig"))}
    dic = {x["eid"]: x for x in csv.DictReader(open(cod, encoding="utf-8-sig"))}
    codif = {x["eid"]: x for x in map(json.loads, open(jsonl, encoding="utf-8"))}
    notas = json.load(open(notas, encoding="utf-8"))
    val = json.load(open(val, encoding="utf-8"))
    usados = [x for x in filas if x["papel"] == "muestra" or x["eid"] in codif]

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for i, linea in enumerate([
        "Supplementary file S5. Full-text coding of the validation sample",
        "",
        "Sample: 40 included studies drawn at random from the 77 open-access studies, stratified by period "
        "(4 from 2016–2019, 10 from 2020–2023, 26 from 2024–2026). When no open full text could be retrieved, "
        "the next replacement of the same period was coded. 2016–2019 had no replacements.",
        "Coder: a large language model (Claude, Anthropic) applying the codebook in Supplementary file S4. "
        "There was no second coder, so no inter-coder agreement is reported.",
        "Every dimension coded as present carries the quotation and page that support it, so that each code can be "
        "audited against the article. Page numbers refer to the PDF page.",
        "Codes for level, design, tradition and object are given only for eligible studies. For borderline "
        "exclusions, the Notes give the codes the study would have received if eligible.",
        "Sheet 'Coding': one row per study. Sheet 'Dictionary accuracy': full-text codes compared with the "
        "dictionary codes for the 25 eligible studies. Sheet 'Summary': eligibility, outcome evaluation and the "
        "three dimensions.",
    ], 1):
        ws.cell(i, 1, linea).font = NEGRITA if i == 1 else NORMAL
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 120

    cab = ["Study", "EID", "DOI", "Year", "Period", "Role", "Language", "Inclusion route in screening",
           "Full text read", "Eligible", "Reason for exclusion", "Level", "Design", "Tradition",
           "Other framework", "Objects of critique", "Outcome evaluation"]
    for d in DIM_EN.values():
        cab += [d, f"{d}: quotation (page)"]
    cab += ["Dictionary codes (title, abstract, keywords)", "Notes"]
    hoja = [cab]
    for x in usados:
        c, d, n = codif.get(x["eid"]), dic[x["eid"]], notas.get(x["eid"], {})
        autor = x["primer_autor"].rsplit(" ", 1)[0]
        fila = [f"{autor} ({x['anio']})", x["eid"], f"https://doi.org/{x['doi']}", int(x["anio"]),
                PERIODO[x["periodo"]], "Sample" if x["papel"] == "muestra" else "Replacement",
                "Spanish" if x["idioma"] == "español" else "English", VIA.get(d["regla"], d["regla"])]
        if not c:
            estado = textos.get(x["eid"], {}).get("estado", "")
            fila.append("No")
            fila += [""] * (len(cab) - len(fila) - 2)
            fila += ["; ".join(en for en, cat in zip(EN_CATEGORIAS, CATEGORIAS) if d[cat] == "1"),
                     "No open full text could be retrieved"
                     + (" (the publisher blocked automated access)" if "403" in estado else "")
                     + "; replaced by the next study of the same period where one was available."]
            hoja.append(fila)
            continue
        fila += ["Yes", "Yes" if c["elig"] else "No", "" if c["elig"] else MOTIVO.get(c["motivo"], c["motivo"])]
        if c["elig"]:
            fila += ["; ".join(ETIQUETAS[k][s] for s in c.get(k, [])) for k in ("niv", "dis", "trad")]
            fila += [n.get("other", ""), "; ".join(ETIQUETAS["obj"][s] for s in c.get("obj", [])),
                     EVALUACION[c["eval"]]]
            for k in DIM_EN:
                cita = c.get("dim", {}).get(k)
                fila += ["Yes", f"“{cita[0]}” ({cita[1]})"] if cita else ["No", ""]
        else:
            fila += [""] * 12
        fila += ["; ".join(en for en, cat in zip(EN_CATEGORIAS, CATEGORIAS) if d[cat] == "1"), n.get("notes", "")]
        hoja.append(fila)
    escribir(wb.create_sheet("Coding"), hoja,
             [22, 20, 30, 6, 11, 12, 9, 30, 9, 9, 22, 22, 26, 28, 34, 34, 24] + [9, 50] * 3 + [44, 70])

    tabla = [["Variable", "Category", "Full text (n)", "Dictionary (n)", "Both", "Dictionary only",
              "Full text only", "Precision", "Sensitivity", "Cohen's κ"]]
    for var, en, cat in zip(VARIABLE, EN_CATEGORIAS, CATEGORIAS):
        r = val["categorias"][cat]
        tabla.append([var, en, r["texto_completo"], r["diccionario"], r["vp"], r["fp"], r["fn"],
                      r["precision"], r["sensibilidad"], r["kappa"]])
    escribir(wb.create_sheet("Dictionary accuracy"), tabla, [18, 40, 12, 13, 8, 14, 13, 10, 11, 10])

    resumen = [["Measure", "Group", "Read", "Eligible"]]
    for p, r in val["elegibilidad_por_periodo"].items():
        resumen.append(["Eligibility by period", p, r["leidos"], r["elegibles"]])
    for v, r in val["elegibilidad_por_via_de_inclusion"].items():
        resumen.append(["Eligibility by inclusion route", VIA.get(v, v), r["leidos"], r["elegibles"]])
    for v, r in val["elegibilidad_por_idioma"].items():
        resumen.append(["Eligibility by language", "Spanish" if v == "español" else "English", r["leidos"], r["elegibles"]])
    for v, r in val["elegibilidad_si_el_diccionario_marca_ia"].items():
        resumen.append(["Eligibility by AI dictionary code", "Coded for AI" if v == "true" else "Not coded for AI",
                        r["leidos"], r["elegibles"]])
    resumen.append(["Eligibility, all", "", val["leidos"], val["elegibles"]])
    resumen += [[], ["Measure", "Category", "Eligible studies (n)", ""]]
    for e in EVALUACION:
        clave = next((k for k in val["evaluacion"] if k.startswith(e[0])), None)
        resumen.append(["Outcome evaluation", e, val["evaluacion"].get(clave, 0) if clave else 0, ""])
    for d, en in zip(DIMENSIONES, DIM_EN.values()):
        resumen.append(["Dimension", en, val["dimensiones"]["frecuencia"][d], ""])
    sigla = {"Txt": "Textual", "Inf": "Infrastructural", "Tra": "Transformative"}
    for comb, nn in val["dimensiones"]["combinaciones"].items():
        resumen.append(["Combination of dimensions", " + ".join(sigla.get(s, s) for s in comb.split("+")), nn, ""])
    resumen += [[], ["Dimension by period", "Period", "Eligible studies (n)", "With the dimension (n)"]]
    for p, r in val["dimensiones"]["por_periodo"].items():
        for d, en in zip(DIMENSIONES, DIM_EN.values()):
            resumen.append([en, p, r["n"], r[d]])
    escribir(wb.create_sheet("Summary"), resumen, [34, 58, 18, 22])
    wb.active = 1
    wb.save(salida)
    print(f"{len(usados)} filas en {salida}")


if __name__ == "__main__":
    main()
